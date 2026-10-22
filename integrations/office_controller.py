"""Controlador de Microsoft Office: Word, Excel, PowerPoint, Outlook."""
from __future__ import annotations

import subprocess
from pathlib import Path
from typing import Any

from core.logger import log

try:
    import win32com.client  # noqa: F401, PLC0415
    WIN32COM_AVAILABLE = True
except ImportError:
    WIN32COM_AVAILABLE = False


class OfficeController:
    """Controla aplicaciones de Microsoft Office.

    Requiere: pip install pywin32
    En Windows, ejecuta después de instalar: python -m pywin32_postinstall -install
    """

    def __init__(self) -> None:
        self.word_app = None
        self.excel_app = None
        self.powerpoint_app = None
        self.outlook_app = None

    # ---- Word ---- #
    def open_word(self, file_path: str | None = None) -> bool:
        """Abre Microsoft Word."""
        try:
            if WIN32COM_AVAILABLE:
                import win32com.client  # noqa: PLC0415
                self.word_app = win32com.client.Dispatch("Word.Application")
                self.word_app.Visible = True

                if file_path:
                    self.word_app.Documents.Open(str(file_path))
                    log.info("[office] Abierto archivo Word: {}", file_path)
                else:
                    self.word_app.Documents.Add()
                    log.info("[office] Nuevo documento Word creado")
                return True
            else:
                # Fallback: abrir directamente la aplicación
                subprocess.Popen("winword.exe" if not file_path else f"winword.exe \"{file_path}\"")
                log.info("[office] Word abierto")
                return True
        except Exception as exc:  # noqa: BLE001
            log.error("[office] Error abriendo Word: {}", exc)
            return False

    def create_word_document(self, content: str, file_path: str) -> bool:
        """Crea un documento de Word con contenido."""
        try:
            if not WIN32COM_AVAILABLE:
                # Fallback usando python-docx
                from docx import Document  # noqa: PLC0415
                doc = Document()
                doc.add_paragraph(content)
                doc.save(file_path)
                log.info("[office] Documento Word creado: {}", file_path)
                return True

            import win32com.client  # noqa: PLC0415
            app = win32com.client.Dispatch("Word.Application")
            app.Visible = True
            doc = app.Documents.Add()
            doc.Content.Text = content
            doc.SaveAs(str(Path(file_path).resolve()))
            log.info("[office] Documento Word creado: {}", file_path)
            return True
        except Exception as exc:  # noqa: BLE001
            log.error("[office] Error creando documento Word: {}", exc)
            return False

    # ---- Excel ---- #
    def open_excel(self, file_path: str | None = None) -> bool:
        """Abre Microsoft Excel."""
        try:
            if WIN32COM_AVAILABLE:
                import win32com.client  # noqa: PLC0415
                self.excel_app = win32com.client.Dispatch("Excel.Application")
                self.excel_app.Visible = True

                if file_path:
                    self.excel_app.Workbooks.Open(str(file_path))
                    log.info("[office] Abierto archivo Excel: {}", file_path)
                else:
                    self.excel_app.Workbooks.Add()
                    log.info("[office] Nuevo libro Excel creado")
                return True
            else:
                subprocess.Popen("excel.exe" if not file_path else f"excel.exe \"{file_path}\"")
                log.info("[office] Excel abierto")
                return True
        except Exception as exc:  # noqa: BLE001
            log.error("[office] Error abriendo Excel: {}", exc)
            return False

    def create_excel_spreadsheet(self, data: list[list[Any]], file_path: str) -> bool:
        """Crea una hoja de Excel con datos.

        Args:
            data: Lista de filas (cada fila es una lista de celdas)
            file_path: Ruta donde guardar el archivo
        """
        try:
            if not WIN32COM_AVAILABLE:
                # Fallback usando openpyxl
                from openpyxl import Workbook  # noqa: PLC0415
                wb = Workbook()
                ws = wb.active
                for row_idx, row_data in enumerate(data, 1):
                    for col_idx, value in enumerate(row_data, 1):
                        ws.cell(row=row_idx, column=col_idx, value=value)
                wb.save(file_path)
                log.info("[office] Hoja Excel creada: {}", file_path)
                return True

            import win32com.client  # noqa: PLC0415
            app = win32com.client.Dispatch("Excel.Application")
            app.Visible = True
            wb = app.Workbooks.Add()
            ws = wb.ActiveSheet

            for row_idx, row_data in enumerate(data, 1):
                for col_idx, value in enumerate(row_data, 1):
                    ws.Cells(row_idx, col_idx).Value = value

            wb.SaveAs(str(Path(file_path).resolve()))
            log.info("[office] Hoja Excel creada: {}", file_path)
            return True
        except Exception as exc:  # noqa: BLE001
            log.error("[office] Error creando Excel: {}", exc)
            return False

    # ---- PowerPoint ---- #
    def open_powerpoint(self, file_path: str | None = None) -> bool:
        """Abre Microsoft PowerPoint."""
        try:
            if WIN32COM_AVAILABLE:
                import win32com.client  # noqa: PLC0415
                self.powerpoint_app = win32com.client.Dispatch("PowerPoint.Application")
                self.powerpoint_app.Visible = True

                if file_path:
                    self.powerpoint_app.Presentations.Open(str(file_path))
                    log.info("[office] Abierto archivo PowerPoint: {}", file_path)
                else:
                    self.powerpoint_app.Presentations.Add()
                    log.info("[office] Nueva presentación PowerPoint creada")
                return True
            else:
                subprocess.Popen("powerpnt.exe" if not file_path else f"powerpnt.exe \"{file_path}\"")
                log.info("[office] PowerPoint abierto")
                return True
        except Exception as exc:  # noqa: BLE001
            log.error("[office] Error abriendo PowerPoint: {}", exc)
            return False

    def create_powerpoint_presentation(self, slides_data: list[dict[str, str]], file_path: str) -> bool:
        """Crea una presentación PowerPoint.

        Args:
            slides_data: Lista de dicts con 'title' y 'content'
            file_path: Ruta donde guardar
        """
        try:
            if not WIN32COM_AVAILABLE:
                # Fallback usando python-pptx
                from pptx import Presentation  # noqa: PLC0415
                prs = Presentation()

                for slide_data in slides_data:
                    slide = prs.slides.add_slide(prs.slide_layouts[1])
                    title = slide.shapes.title
                    body = slide.placeholders[1]

                    title.text = slide_data.get("title", "")
                    text_frame = body.text_frame
                    text_frame.text = slide_data.get("content", "")

                prs.save(file_path)
                log.info("[office] Presentación PowerPoint creada: {}", file_path)
                return True

            import win32com.client  # noqa: PLC0415
            app = win32com.client.Dispatch("PowerPoint.Application")
            app.Visible = True
            prs = app.Presentations.Add()

            for slide_data in slides_data:
                slide_layout = prs.SlideLayouts[1]  # Título y contenido
                slide = prs.Slides.AddSlide(slide_layout)
                title = slide.Shapes.Title
                content = slide.Placeholders[1]

                title.TextFrame.TextRange.Text = slide_data.get("title", "")
                content.TextFrame.TextRange.Text = slide_data.get("content", "")

            prs.SaveAs(str(Path(file_path).resolve()))
            log.info("[office] Presentación PowerPoint creada: {}", file_path)
            return True
        except Exception as exc:  # noqa: BLE001
            log.error("[office] Error creando PowerPoint: {}", exc)
            return False

    # ---- Outlook ---- #
    def open_outlook(self) -> bool:
        """Abre Microsoft Outlook."""
        try:
            if WIN32COM_AVAILABLE:
                import win32com.client  # noqa: PLC0415
                self.outlook_app = win32com.client.Dispatch("Outlook.Application")
                self.outlook_app.Visible = True
                log.info("[office] Outlook abierto")
                return True
            else:
                subprocess.Popen("outlook.exe")
                log.info("[office] Outlook abierto")
                return True
        except Exception as exc:  # noqa: BLE001
            log.error("[office] Error abriendo Outlook: {}", exc)
            return False

    def send_email(self, to: str, subject: str, body: str) -> bool:
        """Envía un correo electrónico via Outlook.

        Args:
            to: Dirección de correo destinatario
            subject: Asunto del correo
            body: Cuerpo del mensaje
        """
        if not WIN32COM_AVAILABLE:
            log.warning("[office] pywin32 no disponible para enviar email via Outlook")
            return False

        try:
            import win32com.client  # noqa: PLC0415
            outlook = win32com.client.Dispatch("Outlook.Application")
            mail = outlook.CreateItem(0)  # 0 = MailItem
            mail.To = to
            mail.Subject = subject
            mail.Body = body
            mail.Send()
            log.info("[office] Email enviado a: {}", to)
            return True
        except Exception as exc:  # noqa: BLE001
            log.error("[office] Error enviando email: {}", exc)
            return False

    def get_outlook_inbox_count(self) -> int:
        """Obtiene el número de emails sin leer en la bandeja de entrada."""
        if not WIN32COM_AVAILABLE:
            return 0

        try:
            import win32com.client  # noqa: PLC0415
            outlook = win32com.client.Dispatch("Outlook.Application")
            inbox = outlook.GetNamespace("MAPI").GetDefaultFolder(6)
            return sum(1 for msg in inbox.Items if msg.UnRead)
        except Exception as exc:  # noqa: BLE001
            log.error("[office] Error obteniendo inbox: {}", exc)
            return 0


office_controller = OfficeController()
